from __future__ import annotations

import argparse
import json
import re
import shutil
import tempfile
import urllib.request
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension
from PIL import Image as PILImage
from PIL import ImageOps, ImageDraw, ImageFont


HEADERS = [
    "Date",
    "Time",
    "Conversation ID",
    "Customer ID",
    "Question",
    "Answer",
    "Model",
    "Unsafe / Wrong Answer",
    "Remarks",
]

MANUAL_HEADERS = [
    "Date",
    "Time",
    "Conversation ID",
    "Customer ID",
    "Title",
    "Reason",
    "Content Type",
]

MONTH_NAMES = {
    1: "January", 2: "February", 3: "March", 4: "April",
    5: "May", 6: "June", 7: "July", 8: "August",
    9: "September", 10: "October", 11: "November", 12: "December",
}

HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"
THIN_GREY = Side(style="thin", color="D9E1F2")
BORDER = Border(left=THIN_GREY, right=THIN_GREY, top=THIN_GREY, bottom=THIN_GREY)

COLUMN_WIDTHS = {
    "A": 13,  # Date
    "B": 11,  # Time
    "C": 39,  # Conversation ID
    "D": 34,  # Customer ID
    "E": 58,  # Question
    "F": 58,  # Answer
    "G": 30,  # Model
    "H": 28,  # Unsafe category
    "I": 62,  # Remarks
}


def load_json(path: Path) -> Dict[str, Any]:
    with path.open("r", encoding="utf-8") as f:
        return json.load(f)


def json_dump(path: Path, obj: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=2)
    tmp.replace(path)


def parse_wib_timestamp(value: Any) -> Tuple[Optional[datetime], Optional[datetime]]:
    """Return (date_datetime, time_datetime) for Excel from MVP1's MM-DD-YYYY HH:MM:SS."""
    if not isinstance(value, str) or not value.strip():
        return None, None
    for fmt in ("%m-%d-%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            dt = datetime.strptime(value.strip(), fmt)
            return datetime(dt.year, dt.month, dt.day), datetime(1899, 12, 30, dt.hour, dt.minute, dt.second)
        except ValueError:
            continue
    return None, None


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def sanitize_remark(value: Any) -> str:
    # Conservative cleanup only: preserve meaning; do not rewrite the evaluator's judgment.
    text = clean_text(value)
    text = re.sub(r"\s+([,.!?;:])", r"\1", text)
    text = re.sub(r"([.!?])\1{2,}", r"\1", text)
    return text[:1200]


def month_sheet_name(month: str) -> str:
    dt = datetime.strptime(month, "%Y-%m")
    return f"{MONTH_NAMES[dt.month]} {dt.year}"


def discover_date_roots(input_root: Path, month: str) -> List[Path]:
    date_re = re.compile(rf"^{re.escape(month)}-(\d{{2}})$")
    roots = [p for p in input_root.iterdir() if p.is_dir() and date_re.match(p.name)] if input_root.exists() else []
    return sorted(roots, key=lambda p: p.name)


def evaluator_result_paths(date_root: Path) -> List[Path]:
    result_dir = date_root / "ai_reviews" / "results"
    if not result_dir.exists():
        return []
    paths = sorted(p for p in result_dir.glob("*.json") if p.is_file())
    master_path = date_root / "master_conversations.json"
    if not master_path.exists():
        return paths
    try:
        master = load_json(master_path)
        active_ids = {
            str(item.get("conversationid"))
            for item in master.get("master_conversations", [])
            if isinstance(item, dict) and item.get("conversationid")
        }
    except Exception:
        return paths
    return [path for path in paths if path.stem in active_ids]


def find_normalized(date_root: Path, conversation_id: str) -> Optional[Path]:
    p = date_root / "normalized" / f"{conversation_id}.json"
    return p if p.exists() else None


def response_for_turn(turn: Dict[str, Any]) -> Tuple[str, List[str], List[str]]:
    """Return text, assistant media URLs, response types for a normalized turn."""
    texts: List[str] = []
    urls: List[str] = []
    types: List[str] = []
    for resp in ((turn.get("assistant") or {}).get("responses") or []):
        if not isinstance(resp, dict):
            continue
        rtype = str(resp.get("type") or "unknown")
        types.append(rtype)
        txt = clean_text(resp.get("text"))
        if txt:
            texts.append(txt)
        for u in resp.get("media_urls") or []:
            if isinstance(u, str) and u:
                urls.append(u)
    return "\n\n".join(texts), urls, types


def find_local_assistant_assets(date_root: Path, conv: Dict[str, Any], turn_index: int) -> List[Path]:
    """Resolve assistant image assets, tolerating stale absolute paths after the dataset was moved."""
    cid = str(conv.get("conversation_id") or "")
    conv_asset_dir = date_root / "assets" / cid
    found: List[Path] = []

    for asset in conv.get("local_assets") or []:
        if not isinstance(asset, dict) or asset.get("error"):
            continue
        if int(asset.get("turn_index") or 0) != turn_index:
            continue
        if str(asset.get("role") or "") != "assistant":
            continue
        raw = asset.get("local_path")
        candidates: List[Path] = []
        if isinstance(raw, str) and raw:
            original = Path(raw)
            candidates.append(original)
            candidates.append(conv_asset_dir / original.name)
            candidates.extend(sorted(conv_asset_dir.glob(original.stem + ".*")))
        for c in candidates:
            if c.exists() and c.is_file() and c not in found:
                found.append(c)
                break
    return found


def best_effort_redownload(url: str, destination: Path) -> Optional[Path]:
    if not url.startswith(("http://", "https://")):
        return None
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=30) as resp, destination.open("wb") as out:
            shutil.copyfileobj(resp, out)
        return destination
    except Exception:
        return None


def fit_image_size(path: Path, max_width: int = 330, max_height: int = 220) -> Tuple[int, int]:
    with PILImage.open(path) as im:
        w, h = im.size
    if w <= 0 or h <= 0:
        return max_width, max_height
    scale = min(max_width / w, max_height / h, 1.0)
    return max(1, int(w * scale)), max(1, int(h * scale))


def compose_images(paths: List[Path], output: Path, max_width: int = 340) -> Optional[Path]:
    """Stack multiple assistant images vertically into one workbook-friendly evidence image."""
    valid: List[PILImage.Image] = []
    try:
        for p in paths:
            try:
                im = PILImage.open(p).convert("RGB")
                target_w = min(max_width, im.width)
                ratio = target_w / im.width if im.width else 1
                target_h = max(1, int(im.height * ratio))
                # Keep each component reasonably bounded so one tall image does not explode row height.
                if target_h > 220:
                    ratio2 = 220 / target_h
                    target_w = max(1, int(target_w * ratio2))
                    target_h = 220
                im = ImageOps.contain(im, (target_w, target_h))
                valid.append(im)
            except Exception:
                continue
        if not valid:
            return None

        gap = 8
        width = max(im.width for im in valid)
        height = sum(im.height for im in valid) + gap * (len(valid) - 1)
        canvas = PILImage.new("RGB", (width, height), "white")
        y = 0
        for im in valid:
            x = (width - im.width) // 2
            canvas.paste(im, (x, y))
            y += im.height + gap
        output.parent.mkdir(parents=True, exist_ok=True)
        canvas.save(output, format="PNG", optimize=True)
        return output
    finally:
        for im in valid:
            try:
                im.close()
            except Exception:
                pass


def question_for_turn(turn: Dict[str, Any]) -> str:
    user = turn.get("user") or {}
    text = clean_text(user.get("text"))
    if text:
        return text
    media = user.get("media_urls") or []
    if media:
        return "[User image input]"
    return "[No user text captured]"


def compact_numbered(items: List[str]) -> str:
    items = [x for x in (clean_text(v) for v in items) if x]
    if not items:
        return ""
    if len(items) == 1:
        return items[0]
    return "\n\n".join(f"[{i}] {v}" for i, v in enumerate(items, start=1))


def collect_month_records(input_root: Path, month: str, media_cache_root: Path) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]], List[str]]:
    unsafe_records: Dict[str, Dict[str, Any]] = {}
    manual_records: Dict[str, Dict[str, Any]] = {}
    warnings: List[str] = []

    for date_root in discover_date_roots(input_root, month):
        for result_path in evaluator_result_paths(date_root):
            try:
                result = load_json(result_path)
            except Exception as e:
                warnings.append(f"Could not read evaluator result {result_path}: {e}")
                continue

            cid = str(result.get("conversation_id") or result_path.stem)
            decision = str(result.get("decision") or "")
            normalized_path = find_normalized(date_root, cid)
            if not normalized_path:
                warnings.append(f"Missing normalized JSON for {cid} ({date_root.name})")
                continue
            try:
                conv = load_json(normalized_path)
            except Exception as e:
                warnings.append(f"Could not read normalized JSON for {cid}: {e}")
                continue

            date_dt, time_dt = parse_wib_timestamp(conv.get("created_at_wib") or result.get("created_at_wib"))
            common = {
                "date": date_dt,
                "time": time_dt,
                "conversation_id": cid,
                "customer_id": clean_text(conv.get("customer_id") or result.get("customer_id")),
                "title": clean_text(conv.get("title") or result.get("title")),
                "date_root": date_root,
            }

            if decision == "REVIEW_REQUIRED":
                types: List[str] = []
                for t in conv.get("turns") or []:
                    user_type = str(((t.get("user") or {}).get("type")) or "")
                    if user_type:
                        types.append(user_type)
                    for resp in ((t.get("assistant") or {}).get("responses") or []):
                        if isinstance(resp, dict) and resp.get("type"):
                            types.append(str(resp.get("type")))
                manual_records[cid] = {
                    **common,
                    "reason": clean_text(result.get("manual_review_reason")) or "Manual review required.",
                    "content_type": ", ".join(sorted(set(types))) or "unknown",
                }
                continue

            if decision != "UNSAFE":
                continue

            unsafe_ids = [str(x) for x in (result.get("unsafe_message_ids") or []) if x]
            unsafe_set = set(unsafe_ids)
            matched_turns: List[Dict[str, Any]] = []
            for t in conv.get("turns") or []:
                if not isinstance(t, dict):
                    continue
                mid = str(t.get("message_id") or "")
                if mid and mid in unsafe_set:
                    matched_turns.append(t)

            if not matched_turns:
                warnings.append(f"No unsafe message IDs could be mapped to turns for {cid}")
                # Keep a visible row instead of silently dropping an evaluated unsafe case.
                unsafe_records[cid] = {
                    **common,
                    "question": "[Unsafe turn mapping unavailable]",
                    "answer_text": "[Unsafe assistant response mapping unavailable]",
                    "models": "",
                    "category": clean_text(result.get("category")),
                    "remark": sanitize_remark(result.get("remark")),
                    "image_paths": [],
                    "image_urls": [],
                }
                continue

            questions: List[str] = []
            answer_texts: List[str] = []
            models: List[str] = []
            image_paths: List[Path] = []
            image_urls: List[str] = []

            for t in matched_turns:
                turn_index = int(t.get("turn_index") or 0)
                questions.append(question_for_turn(t))
                answer_text, urls, _types = response_for_turn(t)
                if answer_text:
                    answer_texts.append(answer_text)
                model = clean_text(t.get("model"))
                if model and model not in models:
                    models.append(model)
                local = find_local_assistant_assets(date_root, conv, turn_index)
                for p in local:
                    if p not in image_paths:
                        image_paths.append(p)
                for u in urls:
                    if u not in image_urls:
                        image_urls.append(u)

            # Best-effort recovery if URLs exist but local assistant evidence is gone.
            if image_urls and not image_paths:
                recovered_dir = media_cache_root / month / cid
                for idx, url in enumerate(image_urls, start=1):
                    parsed_ext = Path(url.split("?", 1)[0]).suffix.lower()
                    if parsed_ext not in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
                        parsed_ext = ".png"
                    dest = recovered_dir / f"assistant_{idx:02d}{parsed_ext}"
                    if dest.exists() or best_effort_redownload(url, dest):
                        image_paths.append(dest)

            unsafe_records[cid] = {
                **common,
                "question": compact_numbered(questions),
                "answer_text": compact_numbered(answer_texts),
                "models": "; ".join(models),
                "category": clean_text(result.get("category")),
                "remark": sanitize_remark(result.get("remark")),
                "image_paths": image_paths,
                "image_urls": image_urls,
            }

    unsafe = sorted(
        unsafe_records.values(),
        key=lambda r: (r.get("date") or datetime.max, r.get("time") or datetime.max, r["conversation_id"]),
    )
    manual = sorted(
        manual_records.values(),
        key=lambda r: (r.get("date") or datetime.max, r.get("time") or datetime.max, r["conversation_id"]),
    )
    return unsafe, manual, warnings


def style_header(ws, headers: List[str]) -> None:
    for col, header in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=header)
        c.fill = PatternFill("solid", fgColor=HEADER_FILL)
        c.font = Font(color=HEADER_FONT, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[1].height = 28


def setup_main_sheet(ws) -> None:
    style_header(ws, HEADERS)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:I1"
    for col, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[col].width = width


def setup_manual_sheet(ws) -> None:
    style_header(ws, MANUAL_HEADERS)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:G1"
    widths = {"A": 13, "B": 11, "C": 39, "D": 34, "E": 48, "F": 65, "G": 24}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def apply_row_style(ws, row: int, cols: int, height: float = 42) -> None:
    ws.row_dimensions[row].height = height
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = BORDER


def write_main_records(ws, records: List[Dict[str, Any]], media_cache_root: Path, month: str) -> int:
    image_count = 0
    for row, rec in enumerate(records, start=2):
        ws.cell(row=row, column=1, value=rec["date"])
        ws.cell(row=row, column=2, value=rec["time"])
        ws.cell(row=row, column=3, value=rec["conversation_id"])
        ws.cell(row=row, column=4, value=rec["customer_id"])
        ws.cell(row=row, column=5, value=rec["question"])
        ws.cell(row=row, column=6, value=rec["answer_text"])
        ws.cell(row=row, column=7, value=rec["models"])
        ws.cell(row=row, column=8, value=rec["category"])
        ws.cell(row=row, column=9, value=rec["remark"])

        ws.cell(row=row, column=1).number_format = "dd-mm-yyyy"
        ws.cell(row=row, column=2).number_format = "hh:mm:ss"
        ws.cell(row=row, column=3).number_format = "@"
        ws.cell(row=row, column=4).number_format = "@"

        image_paths: List[Path] = rec.get("image_paths") or []
        row_height = 52
        if image_paths:
            composite_path = media_cache_root / month / "embedded" / f"{rec['conversation_id']}.png"
            composite = compose_images(image_paths, composite_path)
            if composite:
                w, h = fit_image_size(composite, max_width=330, max_height=320)
                img = XLImage(str(composite))
                img.width = w
                img.height = h
                img.anchor = f"F{row}"
                ws.add_image(img)
                ws.cell(row=row, column=6, value="[AI-generated image embedded]")
                # Excel row height uses points; px * 0.75 is a reasonable approximation.
                row_height = max(row_height, h * 0.78 + 8)
                image_count += len(image_paths)
            elif rec.get("image_urls"):
                ws.cell(row=row, column=6, value="[AI image] " + "\n".join(rec["image_urls"]))
        elif rec.get("image_urls") and not rec.get("answer_text"):
            ws.cell(row=row, column=6, value="[AI image URL] " + "\n".join(rec["image_urls"]))

        apply_row_style(ws, row, 9, row_height)
    return image_count


def write_manual_records(ws, records: List[Dict[str, Any]]) -> None:
    for row, rec in enumerate(records, start=2):
        values = [
            rec["date"], rec["time"], rec["conversation_id"], rec["customer_id"],
            rec["title"], rec["reason"], rec["content_type"],
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        ws.cell(row=row, column=1).number_format = "dd-mm-yyyy"
        ws.cell(row=row, column=2).number_format = "hh:mm:ss"
        ws.cell(row=row, column=3).number_format = "@"
        ws.cell(row=row, column=4).number_format = "@"
        apply_row_style(ws, row, 7, 46)


def build_month_workbook(input_root: Path, output_root: Path, month: str) -> Path:
    media_cache_root = output_root / "_media_cache"
    unsafe, manual, warnings = collect_month_records(input_root, month, media_cache_root)

    wb = Workbook()
    default = wb.active
    wb.remove(default)

    main_name = month_sheet_name(month)
    ws = wb.create_sheet(main_name)
    setup_main_sheet(ws)
    embedded_count = write_main_records(ws, unsafe, media_cache_root, month)

    manual_ws = wb.create_sheet("Manual Review")
    setup_manual_sheet(manual_ws)
    write_manual_records(manual_ws, manual)

    # A compact metadata sheet makes reconciliation/debugging easy while keeping the main tracker clean.
    meta = wb.create_sheet("_Summary")
    meta.sheet_state = "hidden"
    meta_rows = [
        ["Month", month],
        ["Generated at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")],
        ["Unsafe rows", len(unsafe)],
        ["Manual review rows", len(manual)],
        ["Embedded image files", embedded_count],
        ["Warnings", len(warnings)],
    ]
    for r, pair in enumerate(meta_rows, start=1):
        meta.cell(r, 1, pair[0])
        meta.cell(r, 2, pair[1])
    for i, warning in enumerate(warnings, start=len(meta_rows) + 2):
        meta.cell(i, 1, "Warning")
        meta.cell(i, 2, warning)
    meta.column_dimensions["A"].width = 24
    meta.column_dimensions["B"].width = 120

    output_root.mkdir(parents=True, exist_ok=True)
    output = output_root / f"Conversation_Monitoring_{month}.xlsx"
    tmp = output.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    status = "complete"
    canonical_replaced = True
    generated = output
    try:
        tmp.replace(output)
    except PermissionError as e:
        # An open Excel file cannot be atomically replaced on Windows. Preserve it
        # and publish this fully generated workbook under a unique versioned name.
        canonical_replaced = False
        status = "complete_with_locked_copy"
        stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
        generated = output_root / f"Conversation_Monitoring_{month}_{stamp}.xlsx"
        suffix = 2
        while generated.exists():
            generated = output_root / f"Conversation_Monitoring_{month}_{stamp}_{suffix}.xlsx"
            suffix += 1
        try:
            tmp.replace(generated)
        except Exception:
            raise RuntimeError(
                f"Cannot replace locked canonical workbook {output}, and could not save "
                f"versioned fallback {generated}. Temporary workbook remains at {tmp}."
            ) from e
        warnings.append(
            "Canonical workbook was locked; generated a versioned copy instead: "
            f"{generated.name}. Close Excel and rebuild later to update the canonical filename."
        )

    manifest = {
        "month": month,
        "status": status,
        "generated_at": datetime.now().isoformat(),
        "canonical_path": str(output),
        "generated_path": str(generated),
        "canonical_replaced": canonical_replaced,
        "unsafe_rows": len(unsafe),
        "manual_review_rows": len(manual),
        "image_evidence_files": embedded_count,
        "warnings": warnings,
    }
    json_dump(output_root / f"report_build_{month}.json", manifest)

    print("\n=== MONTHLY EXCEL TRACKER — MVP 3 ===")
    print(f"Month                  : {month}")
    print(f"Date folders scanned   : {len(discover_date_roots(input_root, month))}")
    print(f"Unsafe rows            : {len(unsafe)}")
    print(f"Manual review rows     : {len(manual)}")
    print(f"Image evidence files   : {embedded_count}")
    print(f"Warnings               : {len(warnings)}")
    print(f"Report status          : {status}")
    print(f"Workbook               : {generated}")
    if warnings:
        print("\nFirst warnings:")
        for warning in warnings[:10]:
            print(f"  - {warning}")
        if len(warnings) > 10:
            print(f"  ... and {len(warnings)-10} more (see hidden _Summary sheet)")
    return generated


def main() -> int:
    ap = argparse.ArgumentParser(
        description="Dimensions Conversation Monitoring MVP 3 — monthly cumulative Excel tracker"
    )
    group = ap.add_mutually_exclusive_group(required=True)
    group.add_argument("--date", help="A processed date such as 2026-07-01. The whole matching month is rebuilt.")
    group.add_argument("--month", help="Month to rebuild, e.g. 2026-07")
    ap.add_argument("--input-root", default="dimensions_output", help="Root containing YYYY-MM-DD folders")
    ap.add_argument(
        "--output-root",
        default="dimensions_output/monthly_reports",
        help="Destination for monthly workbooks",
    )
    args = ap.parse_args()

    if args.date:
        try:
            dt = datetime.strptime(args.date, "%Y-%m-%d")
        except ValueError:
            print("[ERROR] --date must use YYYY-MM-DD")
            return 2
        month = dt.strftime("%Y-%m")
        date_root = Path(args.input_root).resolve() / args.date
        if not date_root.exists():
            print(f"[ERROR] Date folder not found: {date_root}")
            return 2
    else:
        try:
            datetime.strptime(args.month, "%Y-%m")
        except ValueError:
            print("[ERROR] --month must use YYYY-MM")
            return 2
        month = args.month

    input_root = Path(args.input_root).resolve()
    output_root = Path(args.output_root).resolve()
    roots = discover_date_roots(input_root, month)
    if not roots:
        print(f"[ERROR] No processed date folders found for {month} under {input_root}")
        return 2

    build_month_workbook(input_root, output_root, month)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
